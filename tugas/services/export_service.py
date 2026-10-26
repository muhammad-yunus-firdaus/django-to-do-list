"""
tugas/services/export_service.py
Service layer untuk generator ekspor dokumen: CSV, Excel (OpenPyXL), dan PDF (ReportLab).
"""
import csv
from datetime import date
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles.borders import Border, Side
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from ..models import Kegiatan, AktivitasHarian
from .task_service import get_export_data


def generate_tugas_csv(user):
    """Generate file CSV berisi daftar tugas user."""
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="tugas.csv"'

    writer = csv.writer(response)
    writer.writerow(['Judul', 'Kategori', 'Prioritas', 'Deadline', 'Status'])

    for row in get_export_data(user):
        writer.writerow([
            row["judul"], row["kategori"], row["prioritas"],
            row["deadline"], row["status"]
        ])

    return response


def generate_tugas_excel(user):
    """Generate file Excel (.xlsx) daftar tugas dengan styling."""
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="tugas.xlsx"'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Daftar Tugas"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    border_style = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = ["Judul", "Kategori", "Prioritas", "Deadline", "Status"]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_style

    for row_idx, row in enumerate(get_export_data(user), start=2):
        sheet.append([
            row["judul"], row["kategori"], row["prioritas"],
            row["deadline"], row["status"]
        ])
        for cell in sheet[row_idx]:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = border_style

    for col in sheet.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        sheet.column_dimensions[col[0].column_letter].width = max_len + 2

    workbook.save(response)
    return response


def generate_tugas_pdf(user):
    """Generate file PDF daftar tugas dengan ReportLab."""
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="tugas.pdf"'

    doc = SimpleDocTemplate(response, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    style_normal = styles["Normal"]

    data = [["Judul", "Deskripsi", "Kategori", "Prioritas", "Deadline", "Status"]]

    for row in get_export_data(user):
        data.append([
            Paragraph(row["judul"], style_normal),
            Paragraph(row["deskripsi"], style_normal),
            row["kategori"], row["prioritas"],
            row["deadline"], row["status"],
        ])

    table = Table(data, colWidths=[150, 250, 100, 80, 120, 100])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.gray),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('VALIGN', (0, 1), (-1, -1), 'TOP'),
    ]))

    doc.build([table])
    return response


def generate_kegiatan_pdf(user):
    """Generate file PDF daftar kegiatan & acara."""
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="kegiatan.pdf"'

    doc = SimpleDocTemplate(response, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    style_normal = styles["Normal"]

    data = [["Judul", "Deskripsi", "Tanggal", "Jam Mulai", "Jam Selesai", "Lokasi", "Kategori", "Status"]]

    kegiatan = Kegiatan.objects.filter(user=user).order_by('tanggal', 'jam_mulai')
    for keg in kegiatan:
        data.append([
            Paragraph(keg.judul, style_normal),
            Paragraph(keg.deskripsi or '-', style_normal),
            keg.tanggal.isoformat(),
            keg.jam_mulai.strftime('%H:%M'),
            keg.jam_selesai.strftime('%H:%M') if keg.jam_selesai else '-',
            Paragraph(keg.lokasi or '-', style_normal),
            keg.get_kategori_display(),
            keg.get_status_display(),
        ])

    table = Table(data, colWidths=[120, 180, 80, 60, 60, 100, 100, 100])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.gray),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('VALIGN', (0, 1), (-1, -1), 'TOP'),
    ]))

    doc.build([table])
    return response


def generate_kegiatan_excel(user):
    """Generate file Excel (.xlsx) daftar kegiatan & acara."""
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="kegiatan.xlsx"'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kegiatan & Acara"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    border_style = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = ["Judul", "Deskripsi", "Tanggal", "Jam Mulai", "Jam Selesai", "Lokasi", "Kategori", "Status"]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_style

    kegiatan = Kegiatan.objects.filter(user=user).order_by('tanggal', 'jam_mulai')
    for row_idx, keg in enumerate(kegiatan, start=2):
        sheet.append([
            keg.judul,
            keg.deskripsi or '-',
            keg.tanggal.isoformat(),
            keg.jam_mulai.strftime('%H:%M'),
            keg.jam_selesai.strftime('%H:%M') if keg.jam_selesai else '-',
            keg.lokasi or '-',
            keg.get_kategori_display(),
            keg.get_status_display()
        ])
        for cell in sheet[row_idx]:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = border_style

    for col in sheet.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        sheet.column_dimensions[col[0].column_letter].width = max_len + 2

    workbook.save(response)
    return response


def generate_jadwal_pdf(user, tanggal):
    """Generate file PDF ringkasan jadwal harian A4 portrait."""
    aktivitas_list = AktivitasHarian.objects.filter(
        user=user, tanggal=tanggal
    ).select_related('user').order_by('jam_mulai')

    kegiatan_list = Kegiatan.objects.filter(
        user=user, tanggal=tanggal
    ).select_related('user').order_by('jam_mulai')

    timeline_items = []
    for akt in aktivitas_list:
        timeline_items.append({
            'type': 'Aktivitas Harian',
            'judul': akt.judul,
            'jam_mulai': akt.jam_mulai,
            'jam_selesai': akt.jam_selesai,
            'kategori': 'Habit' if akt.is_habit else 'Rutin',
            'durasi': f"{akt.durasi_menit} mnt",
            'lokasi': '-',
            'status': akt.get_status_display() if hasattr(akt, 'get_status_display') else akt.status,
        })

    for keg in kegiatan_list:
        timeline_items.append({
            'type': 'Kegiatan & Acara',
            'judul': keg.judul,
            'jam_mulai': keg.jam_mulai,
            'jam_selesai': keg.jam_selesai,
            'kategori': keg.get_kategori_display() if hasattr(keg, 'get_kategori_display') else keg.kategori,
            'durasi': '-',
            'lokasi': keg.lokasi or '-',
            'status': keg.get_status_display() if hasattr(keg, 'get_status_display') else keg.status,
        })

    timeline_items.sort(key=lambda x: x['jam_mulai'])

    total_items = len(timeline_items)
    selesai_items = sum(1 for item in timeline_items if item['status'].lower() in ['selesai'])
    progres_persen = round((selesai_items / total_items) * 100) if total_items > 0 else 0

    response = HttpResponse(content_type="application/pdf")
    filename = f"jadwal_{tanggal.isoformat()}.pdf"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    doc = SimpleDocTemplate(response, pagesize=A4, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#4F46E5'),
        spaceAfter=6,
    )
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#475569'),
        spaceAfter=15,
    )
    summary_style = ParagraphStyle(
        'SummaryText',
        parent=styles['Normal'],
        fontSize=11,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=20,
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#0F172A'),
    )

    story = [
        Paragraph("Jadwal Harian & Time-Blocking", title_style),
        Paragraph(f"Tanggal: {tanggal.strftime('%A, %d %B %Y')} | User: {user.username}", meta_style),
        Paragraph(f"Ringkasan: {total_items} Total Agenda | {selesai_items} Selesai ({progres_persen}%)", summary_style),
        Spacer(1, 10),
    ]

    table_data = [["Jam", "Tipe", "Judul", "Kategori", "Durasi", "Lokasi", "Status"]]
    for item in timeline_items:
        jam_str = f"{item['jam_mulai'].strftime('%H:%M')} - {item['jam_selesai'].strftime('%H:%M') if item['jam_selesai'] else '?'}"
        table_data.append([
            jam_str,
            Paragraph(item['type'], cell_style),
            Paragraph(item['judul'], cell_style),
            item['kategori'],
            item['durasi'],
            Paragraph(item['lokasi'], cell_style),
            item['status'].capitalize(),
        ])

    table = Table(table_data, colWidths=[75, 75, 120, 55, 45, 65, 52])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))

    story.append(table)
    doc.build(story)
    return response


def generate_jadwal_excel(user, tanggal):
    """Generate file Excel (.xlsx) agenda jadwal harian."""
    aktivitas_list = AktivitasHarian.objects.filter(
        user=user, tanggal=tanggal
    ).select_related('user').order_by('jam_mulai')

    kegiatan_list = Kegiatan.objects.filter(
        user=user, tanggal=tanggal
    ).select_related('user').order_by('jam_mulai')

    timeline_items = []
    for akt in aktivitas_list:
        timeline_items.append({
            'type': 'Aktivitas Harian',
            'judul': akt.judul,
            'jam_mulai': akt.jam_mulai,
            'jam_selesai': akt.jam_selesai,
            'kategori': 'Habit' if akt.is_habit else 'Rutin',
            'durasi': f"{akt.durasi_menit} mnt",
            'lokasi': '-',
            'status': akt.get_status_display() if hasattr(akt, 'get_status_display') else akt.status,
        })

    for keg in kegiatan_list:
        timeline_items.append({
            'type': 'Kegiatan & Acara',
            'judul': keg.judul,
            'jam_mulai': keg.jam_mulai,
            'jam_selesai': keg.jam_selesai,
            'kategori': keg.get_kategori_display() if hasattr(keg, 'get_kategori_display') else keg.kategori,
            'durasi': '-',
            'lokasi': keg.lokasi or '-',
            'status': keg.get_status_display() if hasattr(keg, 'get_status_display') else keg.status,
        })

    timeline_items.sort(key=lambda x: x['jam_mulai'])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"jadwal_{tanggal.isoformat()}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"Jadwal {tanggal.isoformat()}"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    border_style = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = ["Jam", "Tipe", "Judul", "Kategori", "Durasi", "Lokasi / Link", "Status"]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_style

    for row_idx, item in enumerate(timeline_items, start=2):
        jam_str = f"{item['jam_mulai'].strftime('%H:%M')} - {item['jam_selesai'].strftime('%H:%M') if item['jam_selesai'] else '?'}"
        sheet.append([
            jam_str,
            item['type'],
            item['judul'],
            item['kategori'],
            item['durasi'],
            item['lokasi'],
            item['status']
        ])
        for cell in sheet[row_idx]:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = border_style

    for col in sheet.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        sheet.column_dimensions[col[0].column_letter].width = max_len + 2

    workbook.save(response)
    return response
